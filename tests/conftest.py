import io

import pytest
from docx import Document
from fastapi.testclient import TestClient
from reportlab.pdfgen import canvas

from tenantforge.api import create_app
from tenantforge.auth import Principal
from tenantforge.config import Settings
from tenantforge.processing import chunk, parse
from tenantforge.worker import Worker


@pytest.fixture
def settings(tmp_path):
    return Settings(database_url=f"sqlite:///{tmp_path}/test.db", storage_dir=str(tmp_path / "objects"),
                    daily_token_budget=10_000_000)


@pytest.fixture
def system(settings, monkeypatch):
    # Functional tests use the identical parser synchronously; sandbox execution has separate tests.
    monkeypatch.setattr("tenantforge.worker.parse_isolated",
        lambda data, name, max_chars, max_chunks: chunk(parse(data, name, max_chars), max_chunks=max_chunks))
    app = create_app(settings)
    with TestClient(app) as client:
        svc = app.state.service
        for tenant in ("acme", "beta"):
            svc.db.provision(tenant)
        yield client, svc, Worker(svc.db, svc.storage, svc.providers), app.state.auth


@pytest.fixture
def alice():
    return Principal("acme", "alice", "editor")


def auth_header(auth, tenant="acme", user="alice", role="editor"):
    return {"Authorization": f"Bearer {auth.issue(tenant, user, role)}"}


def upload(system, data=b"Employees receive 25 days of annual leave. The leave year begins in January.",
           name="handbook.txt", tenant="acme", metadata=None):
    import json
    client, svc, worker, auth = system
    response = client.post("/v1/documents", headers=auth_header(auth, tenant),
        files={"file": (name, data)}, data={"metadata": json.dumps(metadata or {"department": "hr"})})
    assert response.status_code == 202, response.text
    worker.once()
    return response.json()


def pdf_bytes(text="Employees receive 25 days of annual leave."):
    out = io.BytesIO()
    pdf = canvas.Canvas(out)
    pdf.drawString(40, 760, text)
    pdf.save()
    return out.getvalue()


def docx_bytes():
    doc = Document()
    doc.add_heading("Leave policy", 1)
    doc.add_paragraph("Employees receive 25 days of annual leave.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Carryover"
    table.rows[0].cells[1].text = "5 days"
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()

