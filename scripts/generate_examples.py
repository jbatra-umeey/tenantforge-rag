"""Generate PDF and DOCX examples with the optional dev dependencies installed."""
from pathlib import Path

from docx import Document
from reportlab.pdfgen import canvas

root = Path(__file__).resolve().parents[1] / "examples"
pdf = canvas.Canvas(str(root / "travel-policy.pdf"))
pdf.drawString(40, 770, "Acme travel policy")
pdf.drawString(40, 740, "Hotel reimbursement is capped at 180 credits per night.")
pdf.drawString(40, 710, "Meals are reimbursed up to 60 credits per day.")
pdf.save()
doc = Document()
doc.add_heading("Acme security policy", 1)
doc.add_paragraph("All customer data is classified as confidential. Production access requires MFA.")
doc.add_paragraph("Security incidents must be reported to the security team within 1 hour.")
doc.save(root / "security-policy.docx")
print("Generated travel-policy.pdf and security-policy.docx")
